"""Build BPMN 2.0 models, diagram images, and the process Word report."""

from pathlib import Path
from xml.etree import ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "procesos_bpmn"
OUT.mkdir(parents=True, exist_ok=True)

# Node: id, BPMN type, label, lane index, column index, optional vertical offset.
from area_models_data import MODELS

NS = {
    "bpmn": "http://www.omg.org/spec/BPMN/20100524/MODEL",
    "bpmndi": "http://www.omg.org/spec/BPMN/20100524/DI",
    "dc": "http://www.omg.org/spec/DD/20100524/DC",
    "di": "http://www.omg.org/spec/DD/20100524/DI",
}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)

FONT = "C:/Windows/Fonts/arial.ttf"
BOLD = "C:/Windows/Fonts/arialbd.ttf"

PROCESS_DETAILS = {
    "10_asistente_de_tickets": {
        "trigger": "Un área necesita crear un ticket y solicita ayuda para clasificarlo.",
        "inputs": "Descripción libre de la necesidad y área de origen.",
        "outputs": "Propuesta revisable de título, clase, categoría, prioridad y área destinataria; ticket solo después de confirmación.",
        "rules": [
            "El asistente usa reglas locales basadas en las funciones de las áreas activas y explica su sugerencia.",
            "Una coincidencia ambigua requiere revisión; TI sigue siendo el destino inicial de las áreas no TI.",
            "La persona puede modificar todos los campos y confirma el envío; el análisis no crea tickets.",
            "SuperAdmin es destino exclusivo de TI para escalamiento.",
        ],
    },
    "09_ti_envia_pedidos_a_areas": {
        "trigger": "TI requiere información, una acción o confirmación de otra área.",
        "inputs": "Área de destino activa, asunto, descripción, clase, categoría y prioridad.",
        "outputs": "Pedido recibido y atendido por el área destinataria, con respuesta e historial; cierre por TI.",
        "rules": [
            "TI elige el destino entre las áreas activas; no puede enviarse un pedido a sí misma.",
            "Solo el área destinataria, TI y SuperAdmin pueden ver el pedido.",
            "El área destinataria toma, atiende o rechaza; TI puede cerrar después de la respuesta.",
            "Si TI elige SuperAdmin, el pedido se registra como Escalado.",
        ],
    },
    "01_coordinacion_entre_areas": {
        "trigger": "Un área detecta una necesidad de documentación, equipo, software, acceso, incidente o servicio.",
        "inputs": "Título, descripción, categoría, prioridad y, cuando corresponda, detalle del pedido.",
        "outputs": "Pedido con código, respuesta del área destinataria, historial y cierre confirmado por el área de origen.",
        "rules": [
            "SIGECO identifica el área de origen y dirige el pedido al área elegida; sin elección, las áreas no TI usan TI como destino inicial.",
            "Cada área consulta los pedidos que origina o recibe; TI y SuperAdmin tienen visibilidad global.",
            "El cierre por el área procede después de Atendido o Rechazado.",
        ],
    },
    "02_administracion": {
        "trigger": "Administración necesita publicar o actualizar un documento, o solicitar un recurso interno.",
        "inputs": "Tipo y contenido del documento; o datos del pedido y prioridad.",
        "outputs": "Documento vigente con versión, o pedido atendido y cerrado.",
        "rules": [
            "La ruta de documentos está autorizada para Administración, TI y SuperAdmin.",
            "Un documento inicia como Borrador y puede recibir nuevas versiones antes o después de publicarse.",
            "Los pedidos de recursos se dirigen a TI y conservan su historial.",
        ],
    },
    "03_ventas": {
        "trigger": "Ventas requiere equipo, software, acceso o servicio para su operación interna.",
        "inputs": "Necesidad concreta, categoría, prioridad y descripción del ticket.",
        "outputs": "Ticket con respuesta, evidencia de atención y cierre de Ventas.",
        "rules": [
            "Ventas solo ve tickets y métricas originados por Ventas.",
            "TI atiende la solicitud y Ventas confirma el cierre después de revisar el resultado.",
            "El flujo no representa clientes, oportunidades ni ventas comerciales.",
        ],
    },
    "04_marketing": {
        "trigger": "Marketing necesita software o un servicio interno para su operación.",
        "inputs": "Descripción del requerimiento, categoría, prioridad y contexto pertinente.",
        "outputs": "Solicitud atendida o rechazada, con historial y cierre del área.",
        "rules": [
            "Marketing consulta únicamente solicitudes originadas por Marketing.",
            "Las evidencias se adjuntan a una transición del ticket y quedan relacionadas con el historial.",
            "El diagrama no representa planificación ni ejecución de campañas.",
        ],
    },
    "05_tecnologia_ti": {
        "trigger": "TI recibe un pedido de otra área en su bandeja global.",
        "inputs": "Pedido, categoría, prioridad, descripción y archivos disponibles.",
        "outputs": "Atención o rechazo registrado; o escalamiento con motivo y resolución de SuperAdmin.",
        "rules": [
            "TI puede tomar pedidos Enviado y pasarlos a EnAtencion.",
            "TI puede marcar Atendido o Rechazado, o escalar con motivo a SuperAdmin.",
            "Después de la atención, el área de origen puede cerrar el pedido.",
        ],
    },
    "06_recursos_humanos": {
        "trigger": "Existe una solicitud de acceso pendiente asociada a Recursos Humanos.",
        "inputs": "Datos de la cuenta, área solicitada y estado pendiente.",
        "outputs": "Cuenta aprobada o solicitud rechazada con estado actualizado.",
        "rules": [
            "Solo el encargado del área o SuperAdmin puede aprobar o rechazar su solicitud.",
            "La cuenta pendiente no puede iniciar sesión hasta recibir aprobación.",
            "La misma regla de aprobación de equipo se aplica en las demás áreas.",
        ],
    },
    "07_contabilidad": {
        "trigger": "Contabilidad necesita documentación, software o acceso para su trabajo interno.",
        "inputs": "Descripción, categoría, prioridad y fecha límite opcional.",
        "outputs": "Solicitud con respuesta, historial y cierre confirmado por Contabilidad.",
        "rules": [
            "Contabilidad ve sus solicitudes y métricas, sin acceso a las otras áreas.",
            "TI atiende; Contabilidad verifica y confirma el cierre.",
            "No se representa facturación, tesorería ni registro contable especializado.",
        ],
    },
    "08_psicologia": {
        "trigger": "Psicología identifica una necesidad interna de acceso, equipo o servicio.",
        "inputs": "Datos del ticket, categoría, prioridad y descripción operacional.",
        "outputs": "Ticket dirigido a TI, respuesta registrada y cierre de Psicología.",
        "rules": [
            "Psicología tiene roles, equipo, panel y visibilidad propios.",
            "Los tickets de Psicología solo son visibles para su área, TI y SuperAdmin.",
            "El proceso no utiliza información clínica ni expedientes personales.",
        ],
    },
}


def geometry(model):
    lane_h = 176
    left, top, gap = 215, 145, 245
    positions = {}
    for node_id, kind, label, lane, col, dy in model["nodes"]:
        w, h = (58, 58) if kind.endswith("Event") else ((80, 80) if kind == "exclusiveGateway" else (205, 80))
        x = left + col * gap + (205 - w) / 2
        y = top + lane * lane_h + (lane_h - h) / 2 + dy
        positions[node_id] = (x, y, w, h)
    last_column = max(node[4] for node in model["nodes"])
    return positions, (left + (last_column + 1) * gap + 50, top + len(model["lanes"]) * lane_h + 80)


def waypoints(a, b, shift=0):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    start = (ax + aw, ay + ah / 2)
    end = (bx, by + bh / 2)
    if end[0] > start[0] + 28:
        mid = (start[0] + end[0]) / 2 + shift
        return [start, (mid, start[1]), (mid, end[1]), end]
    # Loop edge returns above the diagram task row.
    upper = min(start[1], end[1]) - 105
    return [start, (start[0] + 38, start[1]), (start[0] + 38, upper), (end[0] - 30, upper), (end[0] - 30, end[1]), end]


def build_bpmn(model, pos, size):
    b = lambda name: "{%s}%s" % (NS["bpmn"], name)
    d = lambda name: "{%s}%s" % (NS["bpmndi"], name)
    root = ET.Element(b("definitions"), {
        "id": "Definitions_" + model["id"],
        "targetNamespace": "https://nexusadmin.local/bpmn",
        "exporter": "SIGECO BPMN 2.0",
    })
    process_id = "Process_" + model["id"]
    node_by_id = {node[0]: node for node in model["nodes"]}
    process = ET.SubElement(root, b("process"), {"id": process_id, "name": model["title"], "isExecutable": "false"})
    lane_set = ET.SubElement(process, b("laneSet"), {"id": "LaneSet_" + model["id"]})
    for i, lane_name in enumerate(model["lanes"]):
        lane = ET.SubElement(lane_set, b("lane"), {"id": f"Lane_{i}", "name": lane_name})
        for node_id, _, _, lane_index, _, _ in model["nodes"]:
            if lane_index == i:
                ET.SubElement(lane, b("flowNodeRef")).text = node_id
    for node_id, kind, label, *_ in model["nodes"]:
        node = ET.SubElement(process, b(kind), {"id": node_id, "name": label})
        for j, (src, dst, _) in enumerate(model["flows"]):
            if dst == node_id:
                ET.SubElement(node, b("incoming")).text = f"Flow_{j}"
            if src == node_id:
                ET.SubElement(node, b("outgoing")).text = f"Flow_{j}"
    for j, (src, dst, label) in enumerate(model["flows"]):
        ET.SubElement(process, b("sequenceFlow"), {"id": f"Flow_{j}", "sourceRef": src, "targetRef": dst, **({"name": label} if label else {})})

    diagram = ET.SubElement(root, d("BPMNDiagram"), {"id": "Diagram_" + model["id"]})
    plane = ET.SubElement(diagram, d("BPMNPlane"), {"id": "Plane_" + model["id"], "bpmnElement": process_id})
    for i in range(len(model["lanes"])):
        shape = ET.SubElement(plane, d("BPMNShape"), {"id": f"Lane_{i}_di", "bpmnElement": f"Lane_{i}", "isHorizontal": "true"})
        ET.SubElement(shape, "{%s}Bounds" % NS["dc"], {"x": "20", "y": str(145 + i * 176), "width": str(size[0] - 50), "height": "176"})
    for node_id, *_ in model["nodes"]:
        x, y, w, h = pos[node_id]
        shape = ET.SubElement(plane, d("BPMNShape"), {"id": node_id + "_di", "bpmnElement": node_id})
        ET.SubElement(shape, "{%s}Bounds" % NS["dc"], {"x": str(x), "y": str(y), "width": str(w), "height": str(h)})
    for j, (src, dst, _) in enumerate(model["flows"]):
        edge = ET.SubElement(plane, d("BPMNEdge"), {"id": f"Flow_{j}_di", "bpmnElement": f"Flow_{j}"})
        shift = -(node_by_id[src][5] or node_by_id[dst][5]) / 4
        for x, y in waypoints(pos[src], pos[dst], shift):
            ET.SubElement(edge, "{%s}waypoint" % NS["di"], {"x": str(x), "y": str(y)})
    ET.indent(root, space="  ")
    ET.ElementTree(root).write(OUT / (model["id"] + ".bpmn"), encoding="utf-8", xml_declaration=True)


def centered_text(draw, box, label, font, fill, width=22):
    x, y, w, h = box
    lines = []
    current = ""
    for word in label.split():
        candidate = (current + " " + word).strip()
        if current and draw.textbbox((0, 0), candidate, font=font)[2] > w - 5:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    line_h = 23
    yy = y + (h - len(lines) * line_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x + (w - (bbox[2] - bbox[0])) / 2, yy), line, font=font, fill=fill)
        yy += line_h


def build_image(model, pos, size):
    img = Image.new("RGB", size, "#ffffff")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(BOLD, 38)
    lane_font = ImageFont.truetype(BOLD, 22)
    node_font = ImageFont.truetype(FONT, 20)
    small_font = ImageFont.truetype(BOLD, 19)
    node_by_id = {node[0]: node for node in model["nodes"]}
    draw.text((28, 25), model["title"], fill="#12304a", font=title_font)
    draw.text((30, 83), "SIGECO  |  Proceso BPMN 2.0  |  Control operativo por áreas", fill="#597184", font=ImageFont.truetype(FONT, 19))
    for i, name in enumerate(model["lanes"]):
        y = 145 + i * 176
        draw.rectangle((20, y, size[0] - 30, y + 176), fill="#f5f8fb" if i % 2 == 0 else "#edf4f8", outline="#b9c9d4", width=2)
        draw.rectangle((20, y, 195, y + 176), fill="#dfeaf1", outline="#b9c9d4", width=2)
        centered_text(draw, (29, y, 155, 176), name, lane_font, "#12304a", width=15)
    for src, dst, label in model["flows"]:
        shift = -(node_by_id[src][5] or node_by_id[dst][5]) / 4
        pts = waypoints(pos[src], pos[dst], shift)
        draw.line(pts, fill="#657d90", width=4, joint="curve")
        x2, y2 = pts[-1]
        draw.polygon([(x2, y2), (x2 - 13, y2 - 7), (x2 - 13, y2 + 7)], fill="#657d90")
        if label:
            p = pts[-2]
            draw.text((p[0] - 27, pts[-1][1] - 28), label, fill="#184d68", font=small_font)
    for node_id, kind, label, *_ in model["nodes"]:
        x, y, w, h = pos[node_id]
        if kind.endswith("Event"):
            draw.ellipse((x, y, x + w, y + h), fill="#e9f7f3" if kind == "startEvent" else "#fff1ed", outline="#15826d" if kind == "startEvent" else "#bf5a45", width=5 if kind == "endEvent" else 3)
            centered_text(draw, (x - 68, y + h + 3, w + 136, 50), label, node_font, "#233d4d", width=22)
        elif kind == "exclusiveGateway":
            draw.polygon([(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)], fill="#fff1d7", outline="#b98023", width=4)
            draw.line((x + 29, y + 29, x + 51, y + 51), fill="#9b6a1e", width=4)
            draw.line((x + 51, y + 29, x + 29, y + 51), fill="#9b6a1e", width=4)
            centered_text(draw, (x - 60, y + h + 2, w + 120, 50), label, node_font, "#233d4d", width=22)
        else:
            draw.rounded_rectangle((x, y, x + w, y + h), radius=16, fill="#ffffff", outline="#39799a", width=3)
            centered_text(draw, (x + 7, y + 4, w - 14, h - 8), label, node_font, "#17334a")
    img.save(OUT / (model["id"] + ".png"), optimize=True)


def set_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    tc_pr.append(shade)


def add_field(paragraph, field_name):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, code, separate, end):
        run._r.append(item)


def add_labeled(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(label + " ").bold = True
    paragraph.add_run(value)
    return paragraph


def build_word():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = Inches(0.64)
    section.bottom_margin = Inches(0.59)
    section.left_margin = section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor(35, 55, 72)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(23)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(18, 48, 74)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].font.color.rgb = RGBColor(18, 48, 74)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)

    header = section.header.paragraphs[0]
    header.add_run("SIGECO  |  MODELOS DE PROCESOS POR ÁREAS").bold = True
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(80, 105, 123)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("SIGECO  •  ")
    add_field(footer, "PAGE")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(93, 111, 125)

    doc.add_paragraph("Procesos por áreas del Sistema Integrado de Gestión Corporativa SIGECO", style="Title")
    doc.add_paragraph("Documento de modelado BPMN para la presentación final del software de control operativo por áreas.")
    doc.add_heading("Qué representa SIGECO en los diagramas", level=1)
    doc.add_paragraph("El carril SIGECO identifica acciones automáticas del software: analizar la descripción, proponer un destino, asignar el código del pedido, dirigirlo al área elegida, conservar cambios e historial, actualizar estados y limitar la información visible según el área. No representa un departamento ni una persona.")
    doc.add_heading("Alcance del modelado", level=1)
    doc.add_paragraph("El esquema revisado contiene áreas, usuarios y roles, miembros de área, pedidos, detalles, historial, adjuntos, documentos y versiones. Las rutas actuales permiten paneles por área, gestión del equipo y pedidos entre áreas: cada área puede elegir un destino activo y TI queda como destino inicial si no se elige uno. El asistente propone una clasificación para revisión humana. Administración y TI también acceden a gestión documental. Los módulos especializados de Ventas, Marketing, Contabilidad y Psicología no tienen todavía procesos propios más allá de estas funciones compartidas.")
    doc.add_paragraph("Los modelos se entregan como archivos BPMN 2.0 y como imágenes dentro de este Word. Las imágenes se exportaron de los modelos; no son capturas de la aplicación Bizagi.")
    doc.add_heading("Procesos incluidos", level=1)
    for index, model in enumerate(MODELS, 1):
        p = doc.add_paragraph(f"{index:02d}  {model['title']}")
        p.paragraph_format.space_after = Pt(2)

    for index, model in enumerate(MODELS, 1):
        details = PROCESS_DETAILS[model["id"]]
        heading = doc.add_heading(f"{index:02d}  {model['title']}", level=1)
        heading.paragraph_format.page_break_before = True
        doc.add_paragraph(model["description"])
        add_labeled(doc, "Inicio:", details["trigger"])
        add_labeled(doc, "Resultado:", details["outputs"])
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_before = Pt(5)
        picture.paragraph_format.space_after = Pt(2)
        picture.add_run().add_picture(str(OUT / (model["id"] + ".png")), width=Inches(9.45))
        caption = doc.add_paragraph(f"Figura {index}. Diagrama BPMN 2.0 del proceso {model['id']}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.style = styles["Caption"]
        add_labeled(doc, "Alcance funcional:", model["scope"])

        heading = doc.add_heading(f"{index:02d}  Ficha del proceso", level=1)
        heading.paragraph_format.page_break_before = True
        add_labeled(doc, "Proceso:", model["title"])
        add_labeled(doc, "Responsables:", ", ".join(lane for lane in model["lanes"] if lane != "SIGECO"))
        add_labeled(doc, "Entradas:", details["inputs"])
        add_labeled(doc, "Salidas:", details["outputs"])
        doc.add_heading("Secuencia y responsabilidad", level=2)
        tasks = sorted(
            (node for node in model["nodes"] if node[1] in ("task", "exclusiveGateway")),
            key=lambda node: (node[4], node[5]),
        )
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        table.columns[0].width = Inches(0.54)
        table.columns[1].width = Inches(1.56)
        table.columns[2].width = Inches(7.12)
        for cell, title in zip(table.rows[0].cells, ("N.º", "Responsable", "Actividad o decisión")):
            cell.text = title
            set_fill(cell, "17364D")
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(9)
        for step, node in enumerate(tasks, 1):
            row = table.add_row()
            values = (str(step), model["lanes"][node[3]], node[2] + ("  [decisión]" if node[1] == "exclusiveGateway" else ""))
            for cell, value in zip(row.cells, values):
                cell.text = value
                set_fill(cell, "F1F6FA" if step % 2 == 0 else "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.6)
        doc.add_heading("Reglas del proceso", level=2)
        for rule in details["rules"]:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(rule)
        add_labeled(doc, "Sustento técnico:", model["evidence"])

    path = OUT / "Procesos_por_areas_SIGECO.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for model in MODELS:
        pos, size = geometry(model)
        build_bpmn(model, pos, size)
        build_image(model, pos, size)
    print(build_word())
